#!/usr/bin/env python3
"""
PALOMANO LILLE — générateur du dashboard de pilotage.

Lit les 3 fichiers du dossier SUIVI PARC (synchronisés en local par OneDrive),
extrait les données et produit un fichier HTML autonome : Palomano_Dashboard.html

Usage :  python3 build_dashboard.py
"""
import datetime as dt
import json
import os
import re
import sys
from pathlib import Path

import openpyxl

from read_ventes import read_ventes

HERE = Path(__file__).resolve().parent
SRC = Path(os.environ.get(
    "PALOMANO_SUIVI_DIR",
    Path.home() / "Library/CloudStorage/OneDrive-Palomano/PALOMANO LA MAILLERIE - SUIVI PARC",
))
F_SUIVI = SRC / "Objectifs Palomano Lille.xlsx"
F_YANNIKA = SRC / "PALOMANO_Objectifs_Yannika.xlsx"
F_COLOR = SRC / "Colorhisto_PALOMANO_Lille.xlsx"
# Export caisse (facultatif) : le plus récent « palomano-la-maillerie__*.xlsx » du dossier SUIVI PARC,
# ou le fichier désigné par PALOMANO_VENTES
VENTES_GLOB = "palomano-la-maillerie__*.xlsx"
TEMPLATE = HERE / "template.html"
OUT = HERE / "Palomano_Dashboard.html"
OUT_ARTIFACT = HERE / "palomano-cockpit.html"

MONTHS_FR = ["janvier", "février", "mars", "avril", "mai", "juin", "juillet",
             "août", "septembre", "octobre", "novembre", "décembre"]
MONTH_LABELS = ["Janvier", "Février", "Mars", "Avril", "Mai", "Juin", "Juillet",
                "Aout", "Septembre", "Octobre", "Novembre", "Décembre"]


def num(v):
    return float(v) if isinstance(v, (int, float)) else 0.0


def iso(d):
    return d.strftime("%Y-%m-%d")


def r2(x):
    return round(x, 2)


# ───────────────────────── Suivi quotidien ─────────────────────────
def read_suivi():
    wb = openpyxl.load_workbook(F_SUIVI, data_only=True, read_only=True)

    # Journal quotidien (_DATA)
    rows = list(wb["_DATA"].iter_rows(values_only=True))
    daily = []
    for r in rows[1:]:
        d = r[0]
        if not isinstance(d, dt.datetime):
            continue
        daily.append([
            iso(d),
            r2(num(r[6])),   # 1 parc HT
            r2(num(r[9])),   # 2 anniv HT
            r2(num(r[12])),  # 3 scolaire HT
            r2(num(r[15])),  # 4 groupe/centre HT
            r2(num(r[18])),  # 5 snack HT
            r2(num(r[20])),  # 6 boutique HT
            r2(num(r[22])),  # 7 évènement HT
            r2(num(r[24])),  # 8 total HT
            int(num(r[26])),  # 9 entrées parc (tous publics)
            int(num(r[27])),  # 10 capacité
            int(num(r[29])),  # 11 nb anniversaires
            int(num(r[8])),  # 12 entrées parc payantes (hors anniv/groupes)
        ])
    last = max((x[0] for x in daily if x[8] > 0), default=None)
    daily = [x for x in daily if x[0] <= (last or "")]

    # Objectif journalier pondéré (onglets mensuels, colonne « OBJECTIF JOUR »)
    day_obj, day_coef = {}, {}
    for name in wb.sheetnames:
        if not re.match(r"^[A-ZÉa-zéû]+ \d\d$", name):
            continue
        ws = wb[name]
        for r in ws.iter_rows(min_row=9, max_row=40, values_only=True):
            d = r[3] if len(r) > 3 else None
            if isinstance(d, dt.datetime):
                day_obj[iso(d)] = num(r[36])
                day_coef[iso(d)] = num(r[35])

    # Objectifs mensuels calendaires (OBJECTIF ANNUEL : 2025 = 700 k€, 2026 = 650 k€)
    ws = wb["OBJECTIF ANNUEL"]
    obj_cal = {"2025": [0] * 12, "2026": [0] * 12}
    annual = {}
    seen = 0
    for r in ws.iter_rows(values_only=True):
        vals = list(r)
        if "Objectif HT 2025" in vals:
            i = vals.index("Objectif HT 2025")
            annual["2025"] = num(vals[i + 1])
            j = vals.index("Objectif HT 2026")
            annual["2026"] = num(vals[j + 1])
        cells = [v for v in vals if v is not None]
        if cells and cells[0] in MONTH_LABELS and seen < 12 and len(cells) >= 8 \
                and all(isinstance(c, (int, float)) for c in cells[1:4]):
            m = MONTH_LABELS.index(cells[0])
            obj_cal["2025"][m] = num(cells[2])
            obj_cal["2026"][m] = num(cells[6])
            seen += 1

    # Objectifs mensuels par segment (table de l'onglet DASHBOARD, alimentée par Yannika)
    ws = wb["DASHBOARD"]
    obj_seg = {}
    for r in ws.iter_rows(min_row=2, max_row=13, min_col=29, max_col=43, values_only=True):
        if isinstance(r[0], dt.datetime):
            k = r[0].strftime("%Y-%m")
            obj_seg[k] = {"vis": num(r[3]), "anniv": num(r[6]), "b2b": num(r[8]), "ca": num(r[14])}
    wb.close()
    return daily, day_obj, day_coef, obj_cal, annual, obj_seg, last


# ───────────────────────── Objectifs Yannika (BP mai-26 → avr-27) ─────────────────────────
def read_yannika():
    wb = openpyxl.load_workbook(F_YANNIKA, data_only=True)
    ws = wb["Objectifs Yannika"]
    months = []
    for c in range(3, 15):
        h = ws.cell(4, c).value
        months.append(str(h))
    keys = {}
    for r in range(5, 45):
        lab = ws.cell(r, 2).value
        if not isinstance(lab, str):
            continue
        l = lab.strip().lower()
        row = [num(ws.cell(r, c).value) for c in range(3, 15)]
        for k, pat in [("enfants", "entrées enfants"), ("adultes", "entrées adultes"),
                       ("anniv", "ca anniversaires"), ("ecole", "école scolaire"),
                       ("adl", "adl + groupes"), ("palopass", "palo-pass"),
                       ("snackbout", "snacking & boutique"), ("ca", "total ca ht"),
                       ("cap", "capacité max"), ("vis", "visiteurs cibles"),
                       ("ms", "masse salariale totale"), ("staff", "staffme"),
                       ("anim", "animateurs"), ("salaires", "salaires cdi"),
                       ("achats", "sous-total achats"), ("autres", "sous-total autres"),
                       ("charges", "total charges pilotables")]:
            if pat in l and k not in keys:
                keys[k] = row
    annual_target = num(ws["D3"].value)

    # Saisie réelle des charges (onglet Suivi Charges, récap mensuel K13:O24)
    sc = wb["Suivi Charges"]
    charges_real = []
    for r in range(13, 25):
        m = sc.cell(r, 11).value
        charges_real.append({
            "mois": m.strftime("%Y-%m") if isinstance(m, dt.datetime) else str(m),
            "ms": num(sc.cell(r, 12).value), "achats": num(sc.cell(r, 13).value),
            "autres": num(sc.cell(r, 14).value), "total": num(sc.cell(r, 15).value),
        })
    entries = sum(1 for r in range(5, 105) if sc.cell(r, 5).value not in (None, ""))

    # Contrats
    ct = wb["Contrats salariés"]
    contrats = []
    for r in range(6, 13):
        n = ct.cell(r, 2).value
        if n:
            contrats.append({"nom": n, "brut": num(ct.cell(r, 4).value), "coef": num(ct.cell(r, 5).value),
                             "charge": num(ct.cell(r, 6).value),
                             "planning": [num(ct.cell(r + 9, c).value) for c in range(3, 15)]})
    wb.close()
    return {"months": months, "rows": keys, "annual": annual_target,
            "chargesReal": charges_real, "chargesEntries": entries, "contrats": contrats}


# ───────────────────────── Colorhisto / OKR / Avis Google ─────────────────────────
def read_colorhisto():
    wb = openpyxl.load_workbook(F_COLOR, data_only=True)

    ws = wb["Colorhisto hebdo"]
    weeks, wcols = [], []
    for c in range(4, ws.max_column + 1):
        h = ws.cell(4, c).value
        if h:
            weeks.append(str(h).replace("\n", " "))
            wcols.append(c)
    kpis, theme = [], None
    synth = []
    for r in range(5, ws.max_row + 1):
        a, b = ws.cell(r, 1).value, ws.cell(r, 2).value
        if a and not b:
            if str(a).strip().lower().startswith("synth"):
                synth = [ws.cell(r, c).value for c in wcols]
                synth = [round(v, 2) if isinstance(v, (int, float)) else None for v in synth]
            else:
                theme = str(a).strip()
            continue
        if b:
            vals = [ws.cell(r, c).value for c in wcols]
            kpis.append({"theme": theme, "kpi": str(b).strip(), "unit": ws.cell(r, 3).value,
                         "scores": [v if isinstance(v, (int, float)) else None for v in vals]})

    ws = wb["OKR trimestre"]
    okr, cur = [], None
    for r in ws.iter_rows(min_row=1, max_row=ws.max_row, values_only=True):
        a = r[0]
        if isinstance(a, str) and re.match(r"^T\d \d{4}", a):
            cur = {"titre": a.strip(), "rows": []}
            okr.append(cur)
        elif cur is not None and a == "Période":
            cur["debut"] = iso(r[1]) if isinstance(r[1], dt.datetime) else None
            cur["fin"] = iso(r[3]) if isinstance(r[3], dt.datetime) else None
            cur["temps"] = num(r[4])
        elif cur is not None and isinstance(a, str) and isinstance(r[1], (int, float)) and not a.startswith("🎯"):
            cur["rows"].append({"obj": a, "cible": num(r[1]), "reel": num(r[2]),
                                "pct": num(r[3]), "ecart": num(r[4]), "statut": r[5]})
        elif isinstance(a, str) and a.startswith("📖"):
            cur = None

    ws = wb["Avis Google"]
    mois, dist, semaines, releve = [], [], [], None
    mode = None
    for r in ws.iter_rows(values_only=True):
        a = r[0]
        if isinstance(a, str):
            if a.startswith("Relevé du"):
                m = re.search(r"(\d\d/\d\d/\d{4})", a)
                releve = m.group(1) if m else None
            if "PAR MOIS" in a:
                mode = "mois"; continue
            if "DÉTAIL BRUT" in a or "NOTE GLOBALE" in a:
                mode = "dist" if "NOTE GLOBALE" in a else None; continue
            if "PAR SEMAINE" in a:
                mode = "sem"; continue
        if mode == "mois" and isinstance(a, str) and re.match(r"^\d\d/\d{4}", a):
            mois.append({"mois": a, "n": num(r[3]), "note": num(r[4])})
        elif mode == "dist" and isinstance(a, str) and "★" in a:
            dist.append({"etoiles": int(a[0]), "n": num(r[1])})
        elif mode == "sem" and isinstance(r[1], dt.datetime):
            semaines.append({"lib": a, "debut": iso(r[1]), "fin": iso(r[2]), "n": num(r[3]), "note": num(r[4])})
    wb.close()
    return {"weeks": weeks, "kpis": kpis, "synth": synth}, okr, \
        {"releve": releve, "mois": mois, "dist": dist, "semaines": semaines}


def find_ventes():
    env = os.environ.get("PALOMANO_VENTES")
    if env:
        return Path(env) if Path(env).exists() else None
    found = sorted(SRC.glob(VENTES_GLOB), key=lambda f: f.stat().st_mtime)
    return found[-1] if found else None


def main():
    for f in (F_SUIVI, F_YANNIKA, F_COLOR):
        if not f.exists():
            sys.exit(f"Fichier introuvable : {f}")
    daily, day_obj, day_coef, obj_cal, annual, obj_seg, last = read_suivi()
    yannika = read_yannika()
    color, okr, avis = read_colorhisto()
    f_ventes = find_ventes()
    ventes = None
    if f_ventes:
        ventes = read_ventes(f_ventes)
        ventes["file"] = f_ventes.name

    data = {
        "generatedAt": dt.datetime.now().strftime("%Y-%m-%dT%H:%M:%S"),
        "lastDate": last,
        "sources": {f.name: dt.datetime.fromtimestamp(f.stat().st_mtime).strftime("%Y-%m-%dT%H:%M")
                    for f in (F_SUIVI, F_YANNIKA, F_COLOR)},
        "daily": daily,
        "dayObj": {k: r2(v) for k, v in day_obj.items() if v},
        "dayCoef": day_coef,
        "objCal": obj_cal,
        "annual": annual,
        "objSeg": obj_seg,
        "bp": yannika,
        "color": color,
        "okr": okr,
        "avis": avis,
        "ventes": ventes,
    }
    html = TEMPLATE.read_text(encoding="utf-8").replace(
        "/*__DATA__*/null", json.dumps(data, ensure_ascii=False, separators=(",", ":")))
    OUT.write_text(html, encoding="utf-8")

    # Version « Artifact » (claude.ai ajoute lui-même doctype/html/head/body)
    head = re.search(r"<head>(.*?)</head>", html, re.S).group(1)
    body = re.search(r"<body>(.*)</body>", html, re.S).group(1)
    head = re.sub(r'<meta (charset|name="viewport")[^>]*>\s*', "", head)
    OUT_ARTIFACT.write_text(head.strip() + "\n" + body.strip() + "\n", encoding="utf-8")
    print(f"OK → {OUT}  ({len(daily)} jours, dernière saisie {last})")


if __name__ == "__main__":
    main()
